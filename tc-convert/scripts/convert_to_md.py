"""Convert .docx / .xlsx files in a directory to Markdown.

Usage:
  python convert_to_md.py <directory>

Behavior:
  - Converts .docx to sibling .md.
  - Extracts .docx images to <directory>/prodword_pic.
  - Converts .xlsx sheets to Markdown tables.
  - Keeps original files.

Dependencies:
  - python-docx for normal .docx parsing
  - openpyxl for .xlsx parsing
"""

from __future__ import annotations

import os
import re
import sys
import zipfile
from pathlib import Path
from typing import Dict, Iterable, List, Tuple
from xml.etree import ElementTree as ET


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def safe_name(value: str) -> str:
    value = re.sub(r"[\\/:*?\"<>|]+", "_", value)
    value = value.strip().strip(".")
    return value or "image"


def cell_text(cell) -> str:
    return " ".join((cell.text or "").split())


def heading_prefix(style_name: str) -> str | None:
    if not style_name:
        return None
    s = style_name.lower()
    for i in range(1, 7):
        if f"heading {i}" in s or f"标题 {i}" in style_name or s == f"heading{i}":
            return "#" * i
    if "title" in s or style_name.strip() == "标题":
        return "#"
    return None


def image_refs_from_xml(xml_element) -> List[str]:
    refs: List[str] = []
    for blip in xml_element.iter():
        tag = blip.tag.split("}")[-1]
        if tag != "blip":
            continue
        rid = blip.get(f"{R_NS}embed") or blip.get(f"{R_NS}link")
        if rid:
            refs.append(rid)
    return refs


def extract_docx_images(path: Path, pic_dir: Path) -> Dict[str, str]:
    """Extract images from docx and return relationship id -> relative markdown path."""
    pic_dir.mkdir(parents=True, exist_ok=True)
    rel_to_target: Dict[str, str] = {}
    rel_path = "word/_rels/document.xml.rels"

    with zipfile.ZipFile(path) as zf:
        if rel_path in zf.namelist():
            rel_root = ET.fromstring(zf.read(rel_path))
            for rel in rel_root.findall(f"{REL_NS}Relationship"):
                rid = rel.attrib.get("Id")
                target = rel.attrib.get("Target", "")
                rel_type = rel.attrib.get("Type", "")
                if rid and "image" in rel_type and target:
                    rel_to_target[rid] = target

        result: Dict[str, str] = {}
        used_names: set[str] = set()
        for rid, target in rel_to_target.items():
            zip_name = target if target.startswith("word/") else f"word/{target.lstrip('/')}"
            if zip_name not in zf.namelist():
                continue

            original_name = safe_name(Path(target).name)
            stem = safe_name(path.stem)
            candidate = f"{stem}_{original_name}"
            if candidate in used_names:
                suffix = 2
                base = Path(candidate).stem
                ext = Path(candidate).suffix
                while f"{base}_{suffix}{ext}" in used_names:
                    suffix += 1
                candidate = f"{base}_{suffix}{ext}"
            used_names.add(candidate)

            out_path = pic_dir / candidate
            out_path.write_bytes(zf.read(zip_name))
            result[rid] = f"prodword_pic/{candidate}"
        return result


def docx_to_md(path: Path, pic_dir: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    image_map = extract_docx_images(path, pic_dir)
    doc = Document(str(path))
    out: List[str] = []

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            refs = image_refs_from_xml(child)
            image_lines = [f"![图片]({image_map[rid]})" for rid in refs if rid in image_map]
            if not text and not image_lines:
                continue
            hp = heading_prefix(p.style.name if p.style else "")
            if text:
                out.append(f"{hp} {text}" if hp else text)
            out.extend(image_lines)
        elif tag == "tbl":
            t = Table(child, doc)
            rows = t.rows
            if not rows:
                continue
            header = [cell_text(c) for c in rows[0].cells]
            out.append("| " + " | ".join(header) + " |")
            out.append("| " + " | ".join(["---"] * len(header)) + " |")
            for r in rows[1:]:
                cells = [cell_text(c).replace("\n", "<br>") for c in r.cells]
                out.append("| " + " | ".join(cells) + " |")
            out.append("")

    return "\n\n".join(out)


def _el_text(el) -> str:
    parts: List[str] = []
    for node in el.iter():
        tag = node.tag.split("}")[-1]
        if tag == "t":
            parts.append(node.text or "")
        elif tag == "tab":
            parts.append("\t")
        elif tag in ("br", "cr"):
            parts.append("\n")
    return "".join(parts)


def docx_to_md_raw(path: Path, pic_dir: Path) -> str:
    """Fallback parser: read document.xml directly."""
    image_map = extract_docx_images(path, pic_dir)
    with zipfile.ZipFile(str(path)) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    body = root.find(f"{W_NS}body")
    out: List[str] = []
    if body is None:
        return ""

    for child in list(body):
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = _el_text(child).strip()
            refs = image_refs_from_xml(child)
            image_lines = [f"![图片]({image_map[rid]})" for rid in refs if rid in image_map]
            if not text and not image_lines:
                continue
            style = child.find(f"{W_NS}pPr/{W_NS}pStyle")
            prefix = None
            if style is not None:
                val = style.get(f"{W_NS}val", "") or ""
                digits = "".join(ch for ch in val if ch.isdigit())
                if "eading" in val or "标题" in val or (digits and len(digits) == 1):
                    lvl = int(digits) if digits else 1
                    prefix = "#" * min(lvl, 6)
            if text:
                out.append(f"{prefix} {text}" if prefix else text)
            out.extend(image_lines)
        elif tag == "tbl":
            rows = child.findall(f"{W_NS}tr")
            if not rows:
                continue
            md_rows: List[List[str]] = []
            for r in rows:
                cells = r.findall(f"{W_NS}tc")
                md_rows.append([_el_text(c).strip().replace("\n", "<br>") for c in cells])
            width = max(len(r) for r in md_rows)
            md_rows = [r + [""] * (width - len(r)) for r in md_rows]
            out.append("| " + " | ".join(md_rows[0]) + " |")
            out.append("| " + " | ".join(["---"] * width) + " |")
            for r in md_rows[1:]:
                out.append("| " + " | ".join(r) + " |")
            out.append("")

    return "\n\n".join(out)


def xlsx_to_md(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out: List[str] = []
    for ws in wb.worksheets:
        rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            vals = ["" if v is None else str(v).replace("\n", "<br>").strip() for v in row]
            if any(v for v in vals):
                rows.append(vals)
        if not rows:
            continue
        out.append(f"## Sheet: {ws.title}")
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        header = rows[0]
        out.append("| " + " | ".join(header) + " |")
        out.append("| " + " | ".join(["---"] * width) + " |")
        for r in rows[1:]:
            out.append("| " + " | ".join(r) + " |")
        out.append("")
    return "\n\n".join(out)


def iter_inputs(base: Path) -> Iterable[Path]:
    for f in base.rglob("*"):
        if f.is_dir():
            continue
        if "prodword_pic" in f.parts:
            continue
        if f.suffix.lower() not in (".docx", ".xlsx"):
            continue
        if f.name.startswith("~$"):
            continue
        yield f


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: python convert_to_md.py <目录路径>")
        return 2

    base = Path(sys.argv[1]).resolve()
    if not base.exists() or not base.is_dir():
        print(f"ERR 目录不存在: {base}")
        return 1

    pic_dir = base / "prodword_pic"
    converted = 0
    for f in iter_inputs(base):
        md_path = f.with_suffix(".md")
        try:
            if f.suffix.lower() == ".docx":
                try:
                    md = docx_to_md(f, pic_dir)
                except Exception as e1:
                    print(f"    python-docx 失败({e1})，改用 raw 解析")
                    md = docx_to_md_raw(f, pic_dir)
            else:
                md = xlsx_to_md(f)
            md_path.write_text(md, encoding="utf-8")
            print(f"OK  {f.name} -> {md_path.name}  ({len(md)} chars)")
            converted += 1
        except Exception as e:
            print(f"ERR {f.name}: {e}")

    image_count = len(list(pic_dir.glob("*"))) if pic_dir.exists() else 0
    print(f"完成，共转换 {converted} 个文件，抽取图片 {image_count} 张 -> {pic_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
