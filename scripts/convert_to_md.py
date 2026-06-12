"""把目录下的 .docx / .xlsx 转成同名 .md（保留原文件，跳过图片）。
用法: python convert_to_md.py <目录路径>
依赖: python-docx, openpyxl
"""
import sys
import os
from pathlib import Path


def cell_text(cell):
    return " ".join((cell.text or "").split())


def docx_to_md(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out = []

    def heading_prefix(style_name: str):
        if not style_name:
            return None
        s = style_name.lower()
        for i in range(1, 7):
            if f"heading {i}" in s or f"标题 {i}" in style_name or s == f"heading{i}":
                return "#" * i
        if "title" in s or style_name.strip() == "标题":
            return "#"
        return None

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            hp = heading_prefix(p.style.name if p.style else "")
            out.append(f"{hp} {text}" if hp else text)
        elif tag == "tbl":
            t = Table(child, doc)
            rows = t.rows
            if not rows:
                continue
            header = [cell_text(c) for c in rows[0].cells]
            out.append("| " + " | ".join(header) + " |")
            out.append("| " + " | ".join(["---"] * len(header)) + " |")
            for r in rows[1:]:
                cells = [cell_text(c).replace("\n", "<br>") for c in r.cells]
                out.append("| " + " | ".join(cells) + " |")
            out.append("")
    return "\n\n".join(out)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _el_text(el):
    parts = []
    for node in el.iter():
        tag = node.tag.split("}")[-1]
        if tag == "t":
            parts.append(node.text or "")
        elif tag == "tab":
            parts.append("\t")
        elif tag in ("br", "cr"):
            parts.append("\n")
    return "".join(parts)


def docx_to_md_raw(path: Path) -> str:
    """兜底：直接读 word/document.xml，绕开 python-docx 的关系解析。"""
    import zipfile
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(str(path)) as z:
        xml = z.read("word/document.xml")
    root = ET.fromstring(xml)
    body = root.find(f"{W}body")
    out = []
    if body is None:
        return ""
    for child in list(body):
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = _el_text(child).strip()
            if not text:
                continue
            style = child.find(f"{W}pPr/{W}pStyle")
            prefix = None
            if style is not None:
                val = style.get(f"{W}val", "") or ""
                digits = "".join(ch for ch in val if ch.isdigit())
                if "eading" in val or "标题" in val or (digits and len(digits) == 1):
                    lvl = int(digits) if digits else 1
                    prefix = "#" * min(lvl, 6)
            out.append(f"{prefix} {text}" if prefix else text)
        elif tag == "tbl":
            rows = child.findall(f"{W}tr")
            if not rows:
                continue
            md_rows = []
            for r in rows:
                cells = r.findall(f"{W}tc")
                md_rows.append([_el_text(c).strip().replace("\n", "<br>") for c in cells])
            width = max(len(r) for r in md_rows)
            md_rows = [r + [""] * (width - len(r)) for r in md_rows]
            out.append("| " + " | ".join(md_rows[0]) + " |")
            out.append("| " + " | ".join(["---"] * width) + " |")
            for r in md_rows[1:]:
                out.append("| " + " | ".join(r) + " |")
            out.append("")
    return "\n\n".join(out)


def xlsx_to_md(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            vals = ["" if v is None else str(v).replace("\n", "<br>").strip() for v in row]
            if any(v for v in vals):
                rows.append(vals)
        if not rows:
            continue
        out.append(f"## Sheet: {ws.title}")
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        header = rows[0]
        out.append("| " + " | ".join(header) + " |")
        out.append("| " + " | ".join(["---"] * width) + " |")
        for r in rows[1:]:
            out.append("| " + " | ".join(r) + " |")
        out.append("")
    return "\n\n".join(out)


def main():
    if len(sys.argv) < 2:
        print("用法: python convert_to_md.py <目录路径>")
        sys.exit(1)
    base = Path(sys.argv[1])
    converted = 0
    for f in base.rglob("*"):
        if f.suffix.lower() not in (".docx", ".xlsx"):
            continue
        if f.name.startswith("~$"):
            continue
        md_path = f.with_suffix(".md")
        try:
            if f.suffix.lower() == ".docx":
                try:
                    md = docx_to_md(f)
                except Exception as e1:
                    print(f"    python-docx 失败({e1})，改用 raw 解析")
                    md = docx_to_md_raw(f)
            else:
                md = xlsx_to_md(f)
            md_path.write_text(md, encoding="utf-8")
            print(f"OK  {f.name} -> {md_path.name}  ({len(md)} chars)")
            converted += 1
        except Exception as e:
            print(f"ERR {f.name}: {e}")
    print(f"完成，共转换 {converted} 个文件")


if __name__ == "__main__":
    main()
